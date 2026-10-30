"""Independent openpyxl + XML/ZIP/Decimal audit of native evidence workbooks.
Run via uv with openpyxl available: workbook.py OUTPUT_DIRECTORY.
No workbook authoring or scientific recomputation occurs in this checker.
"""
import hashlib
import io
import json
from decimal import Decimal
from pathlib import Path
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

class Number(str):
    pass

DECODER = json.JSONDecoder(parse_int=Number, parse_float=Number)


def nodes(text):
    """Independent recursive scanner retains every scalar's original token."""
    out = {}
    def space(i):
        while i < len(text) and text[i].isspace():
            i += 1
        return i
    def walk(i, path):
        i = space(i)
        first = text[i]
        if first in '{[':
            kind, end = ('object', '}') if first == '{' else ('array', ']')
            i = space(i + 1); count = 0
            while text[i] != end:
                if count:
                    assert text[i] == ','
                    i = space(i + 1)
                if kind == 'object':
                    key, i = DECODER.raw_decode(text, i)
                    assert type(key) is str
                    i = space(i); assert text[i] == ':'
                    i += 1
                    suffix = key.replace('~', '~0').replace('/', '~1')
                else:
                    suffix = str(count)
                i = space(walk(i, path + '/' + suffix)); count += 1
            out[path] = (kind, None, None, count)
            return i + 1
        value, end = DECODER.raw_decode(text, i)
        kind = 'number' if isinstance(value, Number) else 'null' if value is None else 'boolean' if isinstance(value, bool) else 'string'
        assert path not in out
        out[path] = (kind, value, text[i:end], None)
        return end
    end = space(walk(0, ''))
    assert end == len(text)
    return out


def excel_text(value):
    # openpyxl 3.1.5 leaves SpreadsheetML inline-string escape sequences encoded.
    if not isinstance(value, str):
        return value
    return re.sub(r'_x([0-9A-Fa-f]{4})_', lambda m: chr(int(m[1], 16)), value)


def numeric(token):
    d = Decimal(token)
    significant = len(''.join(map(str, d.as_tuple().digits)).strip('0'))
    return not (d.is_zero() and d.is_signed()) and significant <= 15 and (d.is_zero() or Decimal('1e-300') <= abs(d) <= Decimal('1e300'))


def check_sheet(sheet, source):
    expected = nodes(source.decode())
    assert sheet.freeze_panes == 'A2'
    assert sheet.auto_filter.ref == f'A1:E{len(expected)+1}'
    assert sheet.max_row == len(expected) + 1
    seen = set(); textual = 0
    for cells in sheet.iter_rows(min_row=2):
        path, kind, value, exact, children = cells
        pointer = excel_text(path.value)
        assert pointer in expected and pointer not in seen, pointer
        seen.add(pointer)
        source_kind, source_value, token, count = expected[pointer]
        assert kind.value == source_kind
        assert excel_text(exact.value) == token, (pointer, exact.value, token)
        assert children.value == count
        if source_kind in ('object', 'array', 'null'):
            assert value.value is None
        elif source_kind == 'boolean':
            assert value.data_type == 'b' and value.value is source_value
        elif source_kind == 'number' and numeric(token):
            assert value.data_type == 'n'
            assert Decimal(str(value.value)) == Decimal(token), (pointer, value.value, token)
        else:
            assert value.data_type == 's'
            assert excel_text(value.value) == str(source_value), (pointer, value.value, source_value)
            if source_kind == 'number':
                textual += 1
        if token is not None:
            assert exact.data_type == 's'
    assert seen == set(expected)
    return len(seen), textual


def load(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        for name in archive.namelist():
            assert '..' not in name and not name.startswith('/')
            root = ET.fromstring(archive.read(name))
            for el in root.iter():
                assert el.tag.split('}')[-1] not in ('f', 'hyperlink', 'externalLink', 'oleObject')
                assert el.attrib.get('TargetMode') != 'External'
        assert not any('vba' in n.lower() for n in archive.namelist())
    return openpyxl.load_workbook(io.BytesIO(raw), data_only=False)


def audit(folder):
    folder = Path(folder)
    source = Path(__file__).resolve().parents[2] / 'testdata/workbook/precision.json'
    wb = load((folder / 'precision.xlsx').read_bytes())
    count, textual = check_sheet(wb['Precision'], source.read_bytes())
    assert textual == 7
    print(f'Precision fixture: {count} nodes, {textual} exact numeric text cells; types/missingness/escape controls pass.')
    for bundle in sorted(folder.glob('*.zip')):
        with zipfile.ZipFile(bundle) as archive:
            record = json.loads(archive.read('workbook.json'))
            assert record['format'] == 'clonoloom-workbook' and record['version'] == 1
            for item in record['files']:
                data = archive.read(item['path'])
                assert len(data) == item['size'] and hashlib.sha256(data).hexdigest() == item['sha256']
            wb = load(archive.read('results.xlsx'))
            assert wb.sheetnames[0] == 'Read me'
            for i, result in enumerate(record['results'], 1):
                for suffix, sheet_key, stats_key in [('result', 'resultSheet', 'values'), ('manifest', 'provenanceSheet', 'provenance')]:
                    raw = archive.read(f'R{i}.{suffix}.json')
                    count, textual = check_sheet(wb[result[sheet_key]], raw)
                    assert result[stats_key] == {'nodes': count, 'numericText': textual}
                assert hashlib.sha256(archive.read(f'R{i}.result.json')).hexdigest() == result['resultSHA256']
            print(f'{bundle.name}: every recorded node, exact token, cell type and source hash agrees.')

if __name__ == '__main__':
    audit(sys.argv[1])
